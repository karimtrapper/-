"""Выпуск настоящих документов на стенде: POST /api/stand/docs/issue и /upload.

Генератор тот же, что у CRM; LibreOffice в тестах выключен — as_pdf отдаёт DOCX,
его текст и проверяем. Данные синтетические.
"""
import io
import json

import pytest
from docx import Document

import app as appmod
import docgen

RUB_PAY_TO = ('ООО «ЭМ ЭФ КОРПОРЕЙШН» · ИНН 9909726886 · КПП 770387001 · ПАО Сбербанк · '
              'р/с 40807810938720000286 · к/с 30101810400000000225 · БИК 044525225')
PURPOSE = 'Оплата по агентскому договору № SD-9001, НДС не облагается'
GRUSHA = 'TWBgeUo74DehAPgw5cKTdYUTXtJELqwwqn'


def _fields(**over):
    base = {'fio': 'Тестов Иван Петрович', 'fioLat': 'IVAN TESTOV', 'passNo': '75 1234567',
            'passIss': '01.02.2020', 'passOrg': 'МВД 770-001', 'born': '03.04.1985',
            'dev': 'Test Development Co., Ltd.', 'invNo': 'INV-TEST-1', 'invDate': '01.09.2026',
            'object': 'Test Residence, A-101', 'kind': 'Лизхолд',
            'amountThb': '350 000', 'rate': '2,6137', 'amountPay': '914 795',
            'payTo': RUB_PAY_TO, 'purpose': PURPOSE,
            'feeNote': 'Комиссия включена в курс, отдельно не взимается', 'validTill': ''}
    base.update(over)
    return base


def _deal(deal_id, **over):
    d = {'id': deal_id, 'code': f'SD-{deal_id}', 'client': 'Тестов Иван', 'type': 'Оплата недвижимости',
         'kind': 'Лизхолд', 'payType': 'По реквизитам', 'curBase': 'thb', 'step': 's11',
         'isOld': False, 'log': [], 'docParse': {'fields': {
             'client_name_en': 'IVAN TESTOV', 'client_citizenship': 'Россия',
             'project_name': 'Test Residence', 'unit_no': 'A-101'}}}
    d.update(over)
    return d


@pytest.fixture
def stand(monkeypatch):
    monkeypatch.setattr(appmod, 'STAND_MODE', True)
    monkeypatch.setenv('LOCAL_NO_AUTH', '1')

    def role_lookup():
        # как настоящая current_role: закрывает общую scoped-сессию
        appmod.get_session().close()
        return 'operator'
    monkeypatch.setattr(appmod, 'current_role', role_lookup)
    monkeypatch.setattr(docgen, 'to_pdf', lambda raw, timeout=120: None)
    db = appmod.get_session()
    try:
        db.query(appmod.AgreementDoc).delete()
        db.query(appmod.Agreement).delete()
        db.commit()
    finally:
        db.close()

    def put(deals, wallets=None):
        db = appmod.get_session()
        try:
            row = appmod._stand_row(db)
            row.data = json.dumps({'deals': deals, 'convs': [], 'wallets': wallets or []})
            row.version = 1
            db.commit()
        finally:
            db.close()
    with appmod.app.test_client() as client:
        client.put_board = put
        yield client


def _text(client, doc_id):
    r = client.get(f'/api/docs/file/{doc_id}')
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return '\n'.join(parts)


def test_new_client_leasehold_rub_issues_agreement_addendum_invoice(stand):
    stand.put_board([_deal(1)])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 1, 'docFields': _fields()})
    assert r.status_code == 200, r.json
    pack = r.json['pack']
    assert pack['mode'] == 'agreement' and pack['version'] == 1
    assert pack['pair'] == 'RUB_THB' and pack['method'] == 'bank'
    assert pack['purpose'] == PURPOSE
    assert [e['kind'] for e in r.json['issued']] == ['dog', 'app', 'bill']

    saved = stand.get('/api/stand/state').json['data']['deals'][0]
    assert [e['kind'] for e in saved['docsIssued']] == ['dog', 'app', 'bill']
    assert saved['docVersion'] == 1 and saved['docPack']['agreementId'] == pack['agreementId']
    assert 'Выпущен пакет документов' in saved['log'][-1]['text']

    by = {e['kind']: e['docId'] for e in saved['docsIssued']}
    agreement = _text(stand, by['dog'])
    assert 'Тестов Иван Петрович' in agreement and '75 1234567' in agreement
    addendum = _text(stand, by['app'])
    assert '914 795' in addendum and '350 000' in addendum and '2.6137' in addendum
    invoice = _text(stand, by['bill'])
    assert PURPOSE in invoice and '40807810938720000286' in invoice and '914 795' in invoice

    # «Поправить и пересоздать» — тот же договор, новая версия, без второго договора
    again = stand.post('/api/stand/docs/issue', json={'dealId': 1, 'docFields': _fields(amountPay='914 795')})
    assert again.status_code == 200, again.json
    assert again.json['pack']['agreementId'] == pack['agreementId']
    assert again.json['pack']['version'] == 2 and again.json['pack']['mode'] == 'agreement'
    db = appmod.get_session()
    try:
        assert db.query(appmod.Agreement).count() == 1
    finally:
        db.close()


def test_known_client_gets_addendum_and_invoice(stand):
    stand.put_board([_deal(1), _deal(2, isOld=True)])
    first = stand.post('/api/stand/docs/issue', json={'dealId': 1, 'docFields': _fields()})
    assert first.status_code == 200, first.json
    second = stand.post('/api/stand/docs/issue', json={
        'dealId': 2, 'docFields': _fields(amountThb='100 000', rate='2,6', amountPay='260 000')})
    assert second.status_code == 200, second.json
    pack = second.json['pack']
    assert pack['mode'] == 'addendum' and pack['paymentNo'] == 2
    assert pack['agreementId'] == first.json['pack']['agreementId']
    assert [e['kind'] for e in second.json['issued']] == ['app', 'bill']
    invoice = _text(stand, second.json['issued'][1]['docId'])
    assert '260 000' in invoice


def test_known_client_without_agreement_in_base_gets_full_package_with_note(stand):
    stand.put_board([_deal(3, isOld=True)])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 3, 'docFields': _fields()})
    assert r.status_code == 200, r.json
    assert r.json['pack']['mode'] == 'agreement'
    assert 'полный договор' in r.json['pack']['note']


def test_crypto_usdt_thb_uses_wallet_and_no_purpose(stand):
    stand.put_board([_deal(4, payType='Крипта', curBase='usdt', walletId='grusha')])
    fields = _fields(passNo='76 7654321', rate='32,5', amountPay='10 769,23', amountThb='350 000',
                     payTo='USDT · сеть TRC-20 · кошелёк: ' + GRUSHA, purpose='')
    r = stand.post('/api/stand/docs/issue', json={'dealId': 4, 'docFields': fields})
    assert r.status_code == 200, r.json
    pack = r.json['pack']
    assert pack['pair'] == 'USDT_THB' and pack['method'] == 'usdt'
    assert pack['wallet'] == GRUSHA and 'TRC-20' in pack['network'] and pack['purpose'] == ''
    invoice = _text(stand, r.json['issued'][-1]['docId'])
    assert GRUSHA in invoice and '10 769.23' in invoice


def test_missing_fields_return_400_with_stand_labels(stand):
    stand.put_board([_deal(5)])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 5, 'docFields': _fields(passNo='', amountPay='')})
    assert r.status_code == 400
    assert r.json['error'] == 'missing_fields'
    assert set(r.json['fields']) >= {'passNo', 'amountPay'}
    assert 'номер паспорта' in r.json['detail'] and 'сумма клиенту' in r.json['detail']
    db = appmod.get_session()
    try:
        assert db.query(appmod.Agreement).count() == 0
    finally:
        db.close()


def test_rate_not_matching_amounts_is_400_not_500(stand):
    stand.put_board([_deal(6)])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 6, 'docFields': _fields(rate='3,1')})
    assert r.status_code == 400
    assert r.json['error'] == 'rate_mismatch' and 'rate' in r.json['fields']


def test_issue_only_on_document_step(stand):
    stand.put_board([_deal(7, step='s14')])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 7, 'docFields': _fields()})
    assert r.status_code == 409


def test_upload_own_file_replaces_generated(stand):
    stand.put_board([_deal(8)])
    assert stand.post('/api/stand/docs/issue', json={'dealId': 8, 'docFields': _fields()}).status_code == 200
    pdf = b'%PDF-1.4 own contract'
    r = stand.post('/api/stand/docs/upload', data={'dealId': '8', 'kind': 'dog',
                                                   'file': (io.BytesIO(pdf), 'own.pdf')},
                   content_type='multipart/form-data')
    assert r.status_code == 200, r.json
    own = r.json['data']['deals'][0]['issued']['dog']
    assert own['file'] == 'own.pdf'
    assert stand.get(f"/api/docs/file/{own['docId']}").data == pdf
    bad = stand.post('/api/stand/docs/upload', data={'dealId': '8', 'kind': 'dog',
                                                     'file': (io.BytesIO(b'x'), 'x.exe')},
                     content_type='multipart/form-data')
    assert bad.status_code == 400


def test_stand_docs_hidden_outside_stand(monkeypatch):
    monkeypatch.setattr(appmod, 'STAND_MODE', False)
    monkeypatch.setenv('LOCAL_NO_AUTH', '1')
    with appmod.app.test_client() as client:
        assert client.post('/api/stand/docs/issue', json={'dealId': 1}).status_code == 404


def test_zip_contains_current_pack_with_own_file(stand):
    import zipfile
    stand.put_board([_deal(9)])
    assert stand.post('/api/stand/docs/issue', json={'dealId': 9, 'docFields': _fields()}).status_code == 200
    stand.post('/api/stand/docs/upload', data={'dealId': '9', 'kind': 'bill',
                                               'file': (io.BytesIO(b'%PDF-1.4 own bill'), 'bill.pdf')},
               content_type='multipart/form-data')
    r = stand.get('/api/stand/docs/zip?dealId=9')
    assert r.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert len(names) == 3 and 'bill.pdf' in names
    assert any(n.startswith('MF_Agreement_leasehold') for n in names)


def test_leasehold_without_invoice_number_names_the_field(stand):
    stand.put_board([_deal(10)])
    r = stand.post('/api/stand/docs/issue', json={'dealId': 10, 'docFields': _fields(invNo='')})
    assert r.status_code == 400
    assert r.json['fields'] == ['invNo'] and 'номер инвойса' in r.json['detail']
