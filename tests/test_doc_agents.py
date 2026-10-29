# -*- coding: utf-8 -*-
"""Компания-агент в документах: MF Corporation или Flow Sole.

Что ловим:
- выбор агента по счёту прихода (стенд) и по реквизитам оплаты (CRM);
- пакет Flow Sole не содержит ни следа MF (текст, шапка, футер, свойства файла);
- пакет MF не изменился: текст сверяется с эталоном `docgen_mf_golden.json`,
  снятым с генератора ДО появления агентов (коммит b2ca80a);
- платёж на другую компанию — это новый договор, а не допник к чужому.
"""
import io
import json
import os
import re
import zipfile
from datetime import datetime

import pytest
from docx import Document

import app as appmod
import doc_agents
import doc_routes
import docgen
import stand_funding
from tests.test_stand_docs import stand, _deal, _fields, _text  # noqa: F401  (фикстура stand)

WHEN = datetime(2026, 10, 28)
GOLDEN = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'docgen_mf_golden.json')

FS_INN, FS_ACCOUNT, FS_BIK = '9909773251', '40807810819000000261', '044525094'
FS_DETAILS = (f'ИНН: {FS_INN} КПП: 770387001 Расчетный счет: {FS_ACCOUNT} Банк: Московский филиал '
              f'"БАНК СГБ" г. Москва БИК: {FS_BIK} Корр. счет: 30101810245250000094 Валюта счета: Рубли РФ')
FS_RECIPIENT = 'ФЛОУ СОЛЕ КО., ЛТД'
MF_DETAILS = 'Test bank, account 12345'

FIELDS = {
    'client_name_ru': 'Бурова Надежда Васильевна', 'client_name_en': 'Burova Nadezhda',
    'client_citizenship': 'Российской Федерации', 'client_passport_no': '77 2817242',
    'client_passport_issue_date': '14.03.2024', 'client_passport_issued_by': 'МВД 52014',
    'client_passport_expiry_date': '14.03.2034', 'client_birth_date': '02.04.1965',
    'project_name': 'Hythe Condominium By Botanica', 'unit_no': 'B309',
    'property_address': '125/396 Moo. 5 Ratsada Sub-district Mueang Phuket District Phuket, Thailand',
    'lease_term': '60', 'lease_start_date': '20.04.2025',
    'invoice_no': 'BT-B309-007', 'invoice_date': '26.10.2026', 'contract_ref': 'B309',
    'recipient_name': 'Botanica Grand Avenue Co., Ltd.',
    'recipient_bank': 'United Overseas Bank (Thai) Public Company Limited',
    'recipient_account': '602-301-381-6', 'recipient_swift': 'UOVBTHBK',
}


def money(**over):
    base = dict(pair='RUB_THB', payin_currency='RUB', transfer_currency='THB', payin_method='bank',
                total_payin='5 000 000.00', transfer_amount='1 945 500.00', rate='2.570033',
                rate_valid_until='29.10.2026, 23:59', payin_recipient=FS_RECIPIENT,
                payin_recipient_role='Агент', payin_details=FS_DETAILS,
                payment_reference='ОПЛАТА ПО ИНВОЙСУ № BT-B309-007 ОТ 26.10.2026. БЕЗ НДС.')
    base.update(over)
    return base


# ───────────────────────── выбор агента ─────────────────────────

class TestВыборАгентаСтенд:
    def test_сбер_это_mf(self):
        assert appmod._stand_doc_agent({'mode': 'sber'}) == 'mf_corp'

    def test_flowsole_это_flow_sole(self):
        assert appmod._stand_doc_agent({'mode': 'flowsole'}) == 'flow_sole'

    def test_custom_с_реквизитами_flow_sole_это_flow_sole(self):
        account = {'mode': 'custom', 'bank': 'Московский филиал "БАНК СГБ" г. Москва',
                   'account': FS_ACCOUNT, 'bik': FS_BIK, 'correspondent': '30101810245250000094'}
        assert appmod._stand_doc_agent(account) == 'flow_sole'

    def test_custom_с_реквизитами_сбера_это_mf(self):
        account = {'mode': 'custom', 'bank': 'ПАО Сбербанк', 'account': '40807810938720000286',
                   'bik': '044525225', 'correspondent': '30101810400000000225'}
        assert appmod._stand_doc_agent(account) == 'mf_corp'

    def test_custom_чужой_счёт_это_mf(self):
        account = {'mode': 'custom', 'bank': 'Synthetic Bank', 'account': '40702810900000012345',
                   'bik': '044525000', 'correspondent': '30101810000000000000'}
        assert appmod._stand_doc_agent(account) == 'mf_corp'

    def test_счёт_flow_sole_с_чужим_бик_не_узнаётся(self):
        # БИК и счёт узнаются только парой: номер без своего БИК — чужой счёт
        account = {'mode': 'custom', 'bank': 'X', 'account': FS_ACCOUNT, 'bik': '044525000',
                   'correspondent': '30101810245250000094'}
        assert appmod._stand_doc_agent(account) == 'mf_corp'

    def test_usdt_всегда_mf(self):
        assert appmod._stand_doc_agent({'mode': 'flowsole'}, 'usdt') == 'mf_corp'

    @pytest.mark.parametrize('empty', [None, {}, 'text', {'mode': 'wire'}])
    def test_без_счёта_mf(self, empty):
        assert appmod._stand_doc_agent(empty) == 'mf_corp'


class TestВыборАгентаCRM:
    def test_по_инн_flow_sole(self):
        assert doc_agents.agent_for_requisites('', f'ИНН: {FS_INN} Банк: X') == 'flow_sole'

    def test_по_номеру_счёта_flow_sole(self):
        assert doc_agents.agent_for_requisites('', f'р/с {FS_ACCOUNT}') == 'flow_sole'

    def test_получатель_и_реквизиты_вместе(self):
        assert doc_agents.agent_for_money(money()) == 'flow_sole'

    def test_чужие_реквизиты_mf(self):
        assert doc_agents.agent_for_money(money(payin_recipient='ООО «Партнёр»',
                                                payin_details='Банк X, счёт 40702810900000012345')) == 'mf_corp'

    def test_реквизиты_mf_это_mf(self):
        assert doc_agents.agent_for_requisites('', 'ИНН 9909726886, р/с 40807810938720000286') == 'mf_corp'

    def test_реквизиты_двух_компаний_не_угадываем(self):
        text = f'ИНН {FS_INN} и ИНН 9909726886'
        assert doc_agents.agent_for_requisites(text) == 'mf_corp'

    @pytest.mark.parametrize('method', ['usdt', 'cash', 'sbp'])
    def test_не_банковские_способы_всегда_mf(self, method):
        assert doc_agents.agent_for_money(money(payin_method=method)) == 'mf_corp'

    def test_явный_агент_главнее_текста(self):
        assert doc_agents.agent_for_money(money(agent='mf_corp')) == 'mf_corp'

    def test_неизвестный_явный_агент_ошибка(self):
        with pytest.raises(doc_agents.UnknownAgent):
            doc_agents.agent_for_money(money(agent='no_such_company'))


class TestРеестр:
    def test_счета_реестра_совпадают_с_пресетами_прихода(self):
        for key, identity in stand_funding.ACCOUNT_ALIASES.items():
            assert doc_agents.ACCOUNT_REQUISITES[identity]['account'] == key.split(':')[2]

    def test_каждый_известный_счёт_ведёт_к_агенту_реестра(self):
        for identity in stand_funding.ACCOUNT_ALIASES.values():
            assert doc_agents.PAYIN_ACCOUNT_AGENT[identity] in doc_agents.AGENTS

    def test_прочее_ведёт_к_mf(self):
        for identity in (None, 'other:40702810900000012345', 'usdt', 'cash', 'sbp'):
            assert doc_agents.agent_for_account(identity) == 'mf_corp'

    def test_ассеты_каждого_агента_на_месте(self):
        for agent in doc_agents.AGENTS.values():
            for path in (agent.logo_path, agent.signature_path, agent.stamp_path):
                assert os.path.exists(path), path

    def test_ключ_маршрута_mf_прежний(self):
        assert doc_agents.route_key('RUB_THB', 'bank', 'mf_corp') == 'RUB_THB:bank'
        assert doc_agents.route_key('RUB_THB', 'bank') == 'RUB_THB:bank'
        assert doc_agents.route_key('RUB_THB', 'bank', 'flow_sole') == 'RUB_THB:bank:flow_sole'

    def test_старый_ключ_без_агента_это_mf(self):
        assert doc_agents.agent_of_route_key('RUB_THB:bank') == 'mf_corp'
        assert doc_agents.agent_of_route_key('legacy') == 'mf_corp'
        assert doc_agents.agent_of_route_key('RUB_THB:bank:flow_sole') == 'flow_sole'
        assert doc_agents.agent_of_route_key('RUB_THB:bank:unknown_company') == 'mf_corp'

    def test_неизвестный_агент_в_docgen_ошибка(self):
        with pytest.raises(doc_agents.UnknownAgent):
            docgen.build_agreement('leasehold', FIELDS, money(), number='MF-1', when=WHEN, agent='nope')


# ───────────────────────── документы Flow Sole ─────────────────────────

FORBIDDEN = ['MF Corp', 'Katika', 'Sakornnoi', '0835565024547', 'mfcorp', 'RELIABLE PARTNER',
             'Пхукет, Таиланд', 'Choeng Thale', 'Thalang', '+66 991971701', 'mfcorpthai']
REQUIRED = ['Flow Sole Company Limited', '284530232-000', 'Souvahn Phanthavong', 'Вьентьян']
BARE_MF = re.compile(r'(?<![\w-])MF(?![\w-])')


def all_text(data: bytes, raw: bool = True) -> str:
    """Весь текст docx: тело, таблицы, колонтитулы всех секций и (raw) сырой xml частей."""
    z = zipfile.ZipFile(io.BytesIO(data))
    parts = []
    for name in z.namelist():
        if raw and (name.endswith('.xml') or name.endswith('.rels')):
            parts.append(z.read(name).decode('utf8', 'ignore'))
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    for c in docgen._unique_cells(doc):
        texts.append(c.text)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.even_page_header,
                     sec.first_page_footer, sec.even_page_footer):
            texts += [p.text for p in part.paragraphs]
    return '\n'.join(texts + parts)


def build_all(agent, kind='leasehold', **over):
    m = money(**over)
    agreement, _ = docgen.build_agreement(kind, FIELDS, m, number='MF-242-2810-1', when=WHEN, agent=agent)
    full, _ = docgen.build_agreement(kind, FIELDS, m, number='MF-242-2810-1', when=WHEN,
                                     blank_annexes=False, agent=agent)
    addendum = docgen.build_addendum(kind, FIELDS, m, 'MF-242-2810-2', 'MF-242-2810-1', 2, WHEN, agent=agent)
    invoice = docgen.build_commercial_invoice(FIELDS, dict(m, parent_number='MF-242-2810-1'),
                                              'MF-242-2810-2', kind, WHEN, agent=agent)
    short = docgen.build_invoice(FIELDS, m, 'MF-242-2810-2', 'MF-242-2810-1', 'MF-242-2810-2', WHEN, agent=agent)
    return {'agreement': agreement, 'agreement_full': full, 'addendum': addendum,
            'invoice': invoice, 'invoice_short': short}


CASES = [('leasehold', {}), ('freehold', dict(pair='RUB_USD', transfer_currency='USD')),
         ('rental', {}), ('payment', dict(pair='RUB_USDT', transfer_currency='USDT',
                                          payout_network='TRON (TRC-20)', payout_wallet='T' + 'B' * 33)),
         ('leasehold', dict(pair='USDT_THB', payin_currency='USDT', payin_method='usdt',
                            payin_network='TRON (TRC-20)', payin_wallet='T' + 'A' * 33,
                            total_payin='575', transfer_amount='18550', rate='0.030997'))]


class TestДокументыFlowSole:
    @pytest.mark.parametrize('kind,over', CASES)
    def test_нет_следов_mf(self, kind, over):
        for name, data in build_all('flow_sole', kind, **over).items():
            text = all_text(data)
            for bad in FORBIDDEN:
                assert bad not in text, f'{name}: «{bad}»'
            clean = re.sub(r'MF-\d+-\d+-\d+', '', text)
            assert not BARE_MF.search(re.sub(r'<[^>]+>', ' ', clean)), f'{name}: одиночное «MF»'

    @pytest.mark.parametrize('kind,over', CASES)
    def test_реквизиты_flow_sole_на_месте(self, kind, over):
        docs = build_all('flow_sole', kind, **over)
        # договор: все реквизиты; допник и счёт — название, директор, место; налоговый номер
        # стоит в договоре и в строке «Выставил» счёта (как в эталоне Буровой)
        for name in ('agreement', 'agreement_full'):
            text = all_text(docs[name])
            for needed in REQUIRED:
                assert needed in text, f'{name}: «{needed}»'
        addendum = all_text(docs['addendum'])
        for needed in ('Flow Sole Company Limited', 'Souvahn Phanthavong', 'Вьентьян'):
            assert needed in addendum
        assert 'Tel: +856 20 59932483' in all_text(docs['agreement'])
        invoice = all_text(docs['invoice'])
        for needed in ('Flow Sole Company Limited', 'Souvahn Phanthavong', '284530232-000'):
            assert needed in invoice

    def test_директор_в_мужском_роде_и_без_email(self):
        data = build_all('flow_sole')['agreement']
        text = all_text(data, raw=False)
        assert 'действующего на основании' in text and 'действующей' not in text
        assert 'E-mail' not in text and '@' not in text

    def test_подпись_и_печать_flow_sole(self):
        agent = doc_agents.get_agent('flow_sole')
        mf = doc_agents.get_agent('mf_corp')
        data = build_all('flow_sole')['agreement']
        z = zipfile.ZipFile(io.BytesIO(data))
        media = [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
        assert open(agent.signature_path, 'rb').read() in media
        assert open(agent.stamp_path, 'rb').read() in media
        assert open(agent.logo_path, 'rb').read() in media
        for path in (mf.signature_path, mf.stamp_path, mf.logo_path):
            assert open(path, 'rb').read() not in media
        # иконки MF из футера и пустых шапок не остались
        assert all(len(m) > 5000 for m in media)

    def test_роль_и_контакт_в_приложении(self):
        text = all_text(build_all('flow_sole')['addendum'])
        assert 'менеджер Flow Sole / Flow Sole manager' in text
        invoice = all_text(build_all('flow_sole')['invoice'])
        assert 'направить менеджеру Flow Sole / send to the Flow Sole manager' in invoice
        assert 'Комиссия Агента' in invoice

    def test_место_и_дата(self):
        text = all_text(build_all('flow_sole')['agreement'])
        assert 'г. Вьентьян, Лаос / Vientiane, Lao PDR, 28 октября 2026 / 28 October 2026' in text
        text = all_text(build_all('flow_sole')['addendum'])
        assert 'г. Вьентьян, Лаос / Vientiane, Lao PDR, 28 октября 2026' in text

    def test_нумерация_остаётся_mf(self):
        invoice = all_text(build_all('flow_sole')['invoice'])
        assert 'MF-242-2810-2' in invoice

    def test_свойства_файла_без_mf(self):
        data = build_all('flow_sole')['agreement']
        core = zipfile.ZipFile(io.BytesIO(data)).read('docProps/core.xml').decode()
        assert 'MF Corporation' not in core

    def test_агент_берётся_из_money_если_не_передан(self):
        data, _ = docgen.build_agreement('leasehold', FIELDS, money(agent='flow_sole'),
                                         number='MF-1', when=WHEN)
        assert 'Flow Sole Company Limited' in all_text(data)

    def test_check_ловит_чужую_компанию(self):
        data = build_all('mf_corp')['addendum']
        problems = docgen.check(data, agent='flow_sole')
        assert any('MF Corporation Company Limited' in p for p in problems)
        assert docgen.check(data, agent='mf_corp') == []
        assert docgen.check(build_all('flow_sole')['addendum'], agent='flow_sole') == []

    def test_шаблоны_кэшируются_и_не_портят_файлы_mf(self):
        build_all('flow_sole')
        mf_text = all_text(build_all('mf_corp')['agreement'])
        assert 'Katika Sakornnoi' in mf_text and 'Flow Sole' not in mf_text

    def test_нет_висящих_заголовков_разделов(self):
        data = build_all('flow_sole')['agreement']
        doc = Document(io.BytesIO(data))
        rows = doc.tables[0].rows
        heads = [r for r in rows if docgen._is_section_heading(r.cells[-2].text)]
        assert len(heads) == 8
        for row in heads + [r for r in rows if r.cells[-2].text.strip() == 'АГЕНТ / AGENT']:
            assert row._tr.xpath('.//w:keepNext'), row.cells[-2].text
        assert rows[-1]._tr.xpath('.//w:cantSplit')


class TestПакетMFНеИзменился:
    def test_текст_и_картинки_совпадают_с_эталоном(self):
        # эталон снят генератором до появления агентов
        golden = json.load(open(GOLDEN, encoding='utf8'))
        sys_path_tests = os.path.dirname(os.path.abspath(__file__))
        assert os.path.exists(sys_path_tests)
        from datetime import datetime as dt
        from tests.test_docgen import CLIENT

        def dump(data):
            z = zipfile.ZipFile(io.BytesIO(data))
            d = Document(io.BytesIO(data))
            out = [p.text for p in d.paragraphs]
            seen = []
            for t in d.tables:
                for r in t.rows:
                    for c in r.cells:
                        if any(c._tc is s for s in seen):
                            continue
                        seen.append(c._tc)
                        out.append(c.text)
            sec = d.sections[0]
            for part in (sec.header, sec.footer):
                out += [p.text for p in part.paragraphs]
            sizes = sorted(len(z.read(n)) for n in z.namelist() if n.startswith('word/media/'))
            return {'text': out, 'media_sizes': sizes}

        def route(pair, method):
            incoming, outgoing = doc_routes.PAIRS[pair]
            return dict(pair=pair, payin_currency=incoming, transfer_currency=outgoing, payin_method=method,
                        total_payin='575', transfer_amount='18550', rate='0.030997',
                        payin_recipient='MF Corporation Company Limited', payin_recipient_role='Agent',
                        payin_network='TRON (TRC-20)', payin_wallet='T' + 'A' * 33,
                        payin_details='Test bank or SBP details', cash_network='Test cash network',
                        cash_location='Test address', cash_contact='Test cashier',
                        cash_datetime='10.09.2026 15:00 GMT+7', payout_network='Ethereum (ERC-20)',
                        payout_wallet='0x' + '1' * 40, rate_valid_until='10.09.2026 23:59 GMT+7')

        when = dt(2026, 9, 10)
        cases = [('RUB_THB', 'bank', 'leasehold'), ('RUB_THB', 'sbp', 'rental'), ('USDT_THB', 'usdt', 'leasehold'),
                 ('RUB_THB', 'cash', 'leasehold'), ('RUB_USD', 'bank', 'freehold'),
                 ('USDT_USD', 'usdt', 'freehold'), ('RUB_USDT', 'bank', 'payment')]
        got = {}
        for pair, method, kind in cases:
            m = route(pair, method)
            k = f'{pair}:{method}:{kind}'
            got[k + ':agreement'] = dump(docgen.build_agreement(kind, CLIENT, m, when=when)[0])
            got[k + ':addendum'] = dump(docgen.build_addendum(kind, CLIENT, m, 'MF-1', 'MF-0', 1, when=when))
            got[k + ':invoice'] = dump(docgen.build_commercial_invoice(CLIENT, m, 'MF-1', kind, when=when))
        legacy = dict(pair='RUB_THB', payin_method='bank', payin_currency='RUB', total_payin='1000000',
                      transfer_amount='380000', rate='2.63')
        got['legacy:invoice'] = dump(docgen.build_commercial_invoice(CLIENT, legacy, 'MF-1', 'leasehold', when=when))
        assert set(got) == set(golden)
        for key in golden:
            assert got[key] == golden[key], key

    def test_agent_по_умолчанию_mf_и_шапка_шаблона_нетронута(self):
        data, _ = docgen.build_agreement('leasehold', FIELDS, money(payin_recipient='ООО «ЭМ ЭФ КОРПОРЕЙШН»',
                                                                    payin_details=MF_DETAILS), number='MF-1', when=WHEN)
        z = zipfile.ZipFile(io.BytesIO(data))
        original = zipfile.ZipFile(os.path.join(docgen.TEMPLATE_DIR, docgen.TEMPLATES['leasehold']))
        for name in original.namelist():
            if re.fullmatch(r'word/(header|footer)\d\.xml', name) or name.startswith('word/media/'):
                assert z.read(name) == original.read(name), name


# ───────────────────── договор или допник: агент в маршруте ─────────────────────

@pytest.fixture
def crm(monkeypatch):
    monkeypatch.setenv('LOCAL_NO_AUTH', '1')
    appmod.app.config['TESTING'] = True
    monkeypatch.setattr(docgen, 'to_pdf', lambda raw, timeout=120: None)
    db = appmod.get_session()
    try:
        db.query(appmod.AgreementDoc).delete()
        db.query(appmod.Agreement).delete()
        db.commit()
    finally:
        db.close()
    with appmod.app.test_client() as client:
        yield client
    db = appmod.get_session()
    try:
        db.query(appmod.AgreementDoc).delete()
        db.query(appmod.Agreement).delete()
        db.commit()
    finally:
        db.close()


MF_MONEY = dict(payin_recipient='MF Corporation Company Limited', payin_details=MF_DETAILS)
CRM_FIELDS = dict(FIELDS)


def post_agreement(client, **over):
    return client.post('/api/docs/agreements', json={
        'deal_type': 'leasehold', 'fields': CRM_FIELDS, 'money': money(**over)})


def post_payment(client, agreement_id, **over):
    return client.post(f'/api/docs/agreements/{agreement_id}/payment', json={
        'fields': {}, 'money': money(**over)})


def crm_text(client, doc_id):
    response = client.get(f'/api/docs/file/{doc_id}')
    assert response.status_code == 200
    return all_text(response.data)


class TestДоговорИлиДопник:
    def test_у_клиента_договор_mf_платёж_на_flow_sole_новый_договор(self, crm):
        first = post_agreement(crm, **MF_MONEY)
        assert first.status_code == 200, first.json
        mf = first.json['agreement']
        assert mf['agent'] == 'mf_corp' and mf['route_key'] == 'RUB_THB:bank'
        second = post_agreement(crm)
        assert second.status_code == 200, second.json
        fs = second.json['agreement']
        assert fs['id'] != mf['id'] and fs['number'] != mf['number']
        assert fs['agent'] == 'flow_sole' and fs['route_key'] == 'RUB_THB:bank:flow_sole'
        kinds = {d['kind']: d for d in second.json['issued']}
        assert set(kinds) == {'agreement', 'addendum', 'invoice'}
        assert all(d['filename'].startswith('FSC_') for d in kinds.values())
        for doc in kinds.values():
            text = crm_text(crm, doc['id'])
            assert 'Flow Sole Company Limited' in text and 'Katika' not in text
        # файлы договора MF не переписаны
        for doc in first.json['issued']:
            assert doc['filename'].startswith('MF_')
            assert 'Katika' in crm_text(crm, doc['id']) or doc['kind'] == 'invoice'

    def test_и_наоборот_договор_flow_sole_платёж_на_mf(self, crm):
        assert post_agreement(crm).status_code == 200
        second = post_agreement(crm, **MF_MONEY)
        assert second.status_code == 200, second.json
        assert second.json['agreement']['agent'] == 'mf_corp'
        assert second.json['agreement']['route_key'] == 'RUB_THB:bank'

    def test_тот_же_агент_повторный_договор_запрещён(self, crm):
        assert post_agreement(crm).status_code == 200
        again = post_agreement(crm)
        assert again.status_code == 409 and again.json['error'] == 'already_exists'

    def test_допник_к_чужому_договору_не_выпускается(self, crm):
        mf = post_agreement(crm, **MF_MONEY).json['agreement']
        r = post_payment(crm, mf['id'])        # реквизиты Flow Sole в платеже к договору MF
        assert r.status_code == 400 and r.json['error'] == 'invalid_route'
        assert 'Flow Sole Company Limited' in r.json['detail']

    def test_допник_к_договору_flow_sole_от_flow_sole(self, crm):
        fs = post_agreement(crm).json['agreement']
        r = post_payment(crm, fs['id'])
        assert r.status_code == 200, r.json
        assert all(d['filename'].startswith('FSC_') for d in r.json['issued'])
        assert r.json['agreement']['agent'] == 'flow_sole'
        for doc in r.json['issued']:
            assert 'Katika' not in crm_text(crm, doc['id'])

    def test_допник_к_договору_mf_от_mf(self, crm):
        mf = post_agreement(crm, **MF_MONEY).json['agreement']
        r = post_payment(crm, mf['id'], **MF_MONEY)
        assert r.status_code == 200, r.json
        assert all(d['filename'].startswith('MF_') for d in r.json['issued'])

    def test_старая_запись_без_ключа_агента_считается_mf(self, crm):
        mf = post_agreement(crm, **MF_MONEY).json['agreement']
        assert mf['route_key'] == 'RUB_THB:bank'          # у MF ключ прежний, без суффикса
        db = appmod.get_session()
        try:
            row = db.query(appmod.Agreement).get(mf['id'])
            money_json = json.loads(row.money_json)
            money_json.pop('agent', None)                 # как у записи, выпущенной до агентов
            row.money_json = json.dumps(money_json)
            db.commit()
            assert row.to_dict()['agent'] == 'mf_corp'
        finally:
            db.close()
        # платёж на Flow Sole к такому договору — всё равно новый договор, не допник
        assert post_payment(crm, mf['id']).status_code == 400
        assert post_agreement(crm).json['agreement']['agent'] == 'flow_sole'

    def test_неизвестный_агент_отклоняется(self, crm):
        r = post_agreement(crm, agent='nope')
        assert r.status_code == 400 and r.json['error'] == 'invalid_route'

    def test_usdt_маршрут_всегда_mf(self, crm):
        r = post_agreement(crm, pair='USDT_THB', payin_currency='USDT', payin_method='usdt',
                           payin_network='TRON (TRC-20)', payin_wallet='T' + 'A' * 33,
                           payin_details=FS_DETAILS)   # реквизиты Flow Sole в тексте не важны
        assert r.status_code == 200, r.json
        assert r.json['agreement']['agent'] == 'mf_corp'


# ───────────────────────── выпуск пакета на стенде ─────────────────────────

SBER_PAYTO_NAME = 'ООО «ЭМ ЭФ КОРПОРЕЙШН»'


def issue(stand, deal_id, **fields):
    return stand.post('/api/stand/docs/issue', json={'dealId': deal_id, 'docFields': _fields(**fields)})


def board_deal(stand, deal_id):
    return next(d for d in stand.get('/api/stand/state').json['data']['deals'] if d['id'] == deal_id)


class TestСтендАгент:
    def test_flowsole_выпускает_flow_sole(self, stand):
        stand.put_board([_deal(1001, rubReceivingAccount={'mode': 'flowsole'})])
        r = issue(stand, 1001)
        assert r.status_code == 200, r.json
        deal = board_deal(stand, 1001)
        assert deal['issuedDocAgent'] == 'flow_sole'
        assert deal['docPack']['agent'] == 'flow_sole'
        assert all(e['agent'] == 'flow_sole' for e in deal['docsIssued'])
        assert all(e['file'].startswith('FSC_') for e in deal['docsIssued'])
        for entry in r.json['issued']:
            text = _text(stand, entry['docId'])
            assert 'Flow Sole Company Limited' in text
            assert 'Katika' not in text and 'MF Corporation' not in text

    def test_sber_выпускает_mf(self, stand):
        stand.put_board([_deal(1002, rubReceivingAccount={'mode': 'sber'})])
        r = issue(stand, 1002)
        assert r.status_code == 200, r.json
        deal = board_deal(stand, 1002)
        assert deal['issuedDocAgent'] == 'mf_corp' and deal['docPack']['agent'] == 'mf_corp'
        assert all(e['file'].startswith('MF_') for e in deal['docsIssued'])
        assert 'Katika Sakornnoi' in _text(stand, next(e['docId'] for e in r.json['issued'] if e['kind'] == 'dog'))

    def test_без_rubReceivingAccount_mf(self, stand):
        stand.put_board([_deal(1003)])
        assert issue(stand, 1003).status_code == 200
        assert board_deal(stand, 1003)['issuedDocAgent'] == 'mf_corp'

    def test_custom_с_реквизитами_flow_sole(self, stand):
        account = {'mode': 'custom', 'bank': 'Московский филиал "БАНК СГБ" г. Москва',
                   'account': FS_ACCOUNT, 'correspondent': '30101810245250000094', 'bik': FS_BIK}
        stand.put_board([_deal(1004, rubReceivingAccount=account)])
        r = issue(stand, 1004)
        assert r.status_code == 200, r.json
        deal = board_deal(stand, 1004)
        assert deal['issuedDocAgent'] == 'flow_sole'
        # получатель не назван → это сам Flow Sole, а не ЭМ ЭФ с чужим счётом
        assert deal['issuedDocPayTo'].startswith('ФЛОУ СОЛЕ КО., ЛТД · ИНН 9909773251 · КПП 770387001 · ')
        assert 'ЭМ ЭФ' not in deal['issuedDocPayTo']
        invoice = _text(stand, next(e['docId'] for e in r.json['issued'] if e['kind'] == 'bill'))
        assert 'ЭМ ЭФ' not in invoice and FS_ACCOUNT in invoice

    def test_custom_чужой_счёт_mf(self, stand):
        account = {'mode': 'custom', 'bank': 'Synthetic Bank', 'account': '40702810900000012345',
                   'correspondent': '30101810000000000000', 'bik': '044525225'}
        stand.put_board([_deal(1005, rubReceivingAccount=account)])
        assert issue(stand, 1005).status_code == 200
        assert board_deal(stand, 1005)['issuedDocAgent'] == 'mf_corp'

    def test_usdt_mf_даже_при_счёте_flow_sole_в_сделке(self, stand):
        wallet = 'TKkeEVf2zySaWTLyX2qPwvi6kcdHRuPxkJ'
        stand.put_board([_deal(1006, payType='Крипта', curBase='usdt', walletId='grusha',
                               rubReceivingAccount={'mode': 'flowsole'})],
                        wallets=[{'id': 'grusha', 'addr': wallet}])
        r = issue(stand, 1006, payTo=f'USDT · сеть TRC-20 · кошелёк: {wallet}', purpose='',
                  rate='32,5', amountPay='10 769,23', amountThb='350 000')
        assert r.status_code == 200, r.json
        deal = board_deal(stand, 1006)
        assert deal['issuedDocAgent'] == 'mf_corp'
        assert all(e['file'].startswith('MF_') for e in deal['docsIssued'])

    def test_клиент_с_договором_mf_и_платёж_на_flow_sole_новый_договор(self, stand):
        stand.put_board([_deal(1010, rubReceivingAccount={'mode': 'sber'}),
                         _deal(1011, rubReceivingAccount={'mode': 'flowsole'}),
                         _deal(1012, rubReceivingAccount={'mode': 'sber'}),
                         _deal(1013, rubReceivingAccount={'mode': 'flowsole'})])
        first = issue(stand, 1010)
        assert first.status_code == 200 and board_deal(stand, 1010)['docPack']['mode'] == 'agreement'
        second = issue(stand, 1011)
        assert second.status_code == 200, second.json
        pack = board_deal(stand, 1011)['docPack']
        assert pack['mode'] == 'agreement', 'платёж на Flow Sole оформлен допником к договору MF'
        assert pack['agent'] == 'flow_sole'
        assert pack['agreementNumber'] != board_deal(stand, 1010)['docPack']['agreementNumber']
        kinds = [e['kind'] for e in board_deal(stand, 1011)['docsIssued']]
        assert 'dog' in kinds
        # следующий платёж на тот же счёт — допник к договору ТОЙ ЖЕ компании
        third = issue(stand, 1012)
        assert third.status_code == 200 and board_deal(stand, 1012)['docPack']['mode'] == 'addendum'
        assert (board_deal(stand, 1012)['docPack']['agreementNumber']
                == board_deal(stand, 1010)['docPack']['agreementNumber'])
        fourth = issue(stand, 1013)
        assert fourth.status_code == 200 and board_deal(stand, 1013)['docPack']['mode'] == 'addendum'
        assert (board_deal(stand, 1013)['docPack']['agreementNumber']
                == board_deal(stand, 1011)['docPack']['agreementNumber'])

    def test_перевыпуск_сохраняет_агента(self, stand):
        stand.put_board([_deal(1020, rubReceivingAccount={'mode': 'flowsole'})])
        assert issue(stand, 1020).status_code == 200
        first_number = board_deal(stand, 1020)['docPack']['agreementNumber']
        again = issue(stand, 1020, amountPay='915 000')
        assert again.status_code == 200, again.json
        pack = board_deal(stand, 1020)['docPack']
        assert pack['agent'] == 'flow_sole' and pack['agreementNumber'] == first_number

    def test_смена_счёта_после_выпуска_выпускает_документы_новой_компании(self, stand):
        stand.put_board([_deal(1030, rubReceivingAccount={'mode': 'sber'})])
        assert issue(stand, 1030).status_code == 200
        old_files = [e['file'] for e in board_deal(stand, 1030)['docsIssued']]
        old_number = board_deal(stand, 1030)['docPack']['agreementNumber']
        # оператор поправил счёт на Flow Sole и пересоздаёт пакет
        db = appmod.get_session()
        try:
            row = appmod._stand_row(db)
            state = json.loads(row.data)
            state['deals'][0]['rubReceivingAccount'] = {'mode': 'flowsole'}
            row.data = json.dumps(state)
            row.version = (row.version or 0) + 1
            db.commit()
        finally:
            db.close()
        r = issue(stand, 1030)
        assert r.status_code == 200, r.json
        deal = board_deal(stand, 1030)
        assert deal['issuedDocAgent'] == 'flow_sole'
        assert deal['docPack']['agent'] == 'flow_sole'
        assert deal['docPack']['mode'] == 'agreement' and deal['docPack']['agreementNumber'] != old_number
        assert 'Реквизиты оплаты сменили компанию-агента' in deal['docPack']['note']
        # прежние файлы MF остались в списке выпущенного
        assert [e['file'] for e in deal['docsIssued']][:len(old_files)] == old_files
        db = appmod.get_session()
        try:
            assert db.query(appmod.Agreement).filter(appmod.Agreement.route_key == 'RUB_THB:bank').count() == 1
            assert db.query(appmod.Agreement).filter(
                appmod.Agreement.route_key == 'RUB_THB:bank:flow_sole').count() == 1
        finally:
            db.close()
